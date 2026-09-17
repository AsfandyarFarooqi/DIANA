import pytest
from docx import Document

from diana.documents import export_document, parse_blocks

SAMPLE = """## Summary

**To:** Ms. Khan
**From:** Ali

| Item | Qty |
|------|-----|
| Laptop | 2 |

- one
- two

1. first

---

```
code()
```

End & <done>."""


def test_parse_blocks_recognises_structure():
    blocks = parse_blocks(SAMPLE)
    assert [kind for kind, _ in blocks] == ["h", "p", "table", "ul", "ul", "ol", "hr", "code", "p"]
    assert blocks[1][1] == "**To:** Ms. Khan\n**From:** Ali"
    assert blocks[2][1] == [["Item", "Qty"], ["Laptop", "2"]]


@pytest.mark.parametrize("fmt, magic", [("docx", b"PK"), ("pdf", b"%PDF"), ("md", b"# Report: Q3 & Beyond")])
def test_export_formats(fmt, magic):
    buf, mime, filename = export_document("Report: Q3 & Beyond", SAMPLE, fmt)
    assert buf.read().startswith(magic)
    assert filename == f"Report_Q3_Beyond.{fmt}"
    assert mime


def test_docx_keeps_line_breaks_and_tables():
    buf, _, _ = export_document("Letter", SAMPLE, "docx")
    doc = Document(buf)
    assert "To: Ms. Khan\nFrom: Ali" in [p.text for p in doc.paragraphs]
    assert doc.tables[0].cell(1, 0).text == "Laptop"


def test_export_rejects_bad_input():
    with pytest.raises(ValueError):
        export_document("Empty", "   ", "pdf")
    with pytest.raises(ValueError):
        export_document("Unknown", "text", "exe")
