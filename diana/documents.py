"""Render Markdown-ish document content to DOCX / PDF / MD, fully in memory."""
import io
import os
import re

from docx import Document
from docx.shared import Pt

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")


def parse_blocks(content):
    """Split Markdown into simple blocks: (kind, payload)."""
    blocks, para, table, code = [], [], [], None

    def flush():
        if para:
            blocks.append(("p", "\n".join(para)))  # keep line breaks: letters rely on them
            para.clear()

    def flush_table():
        if not table:
            return
        rows = [[c.strip() for c in r.strip("|").split("|")] for r in table]
        if len(rows) >= 2 and all(re.fullmatch(r":?-{2,}:?", c) for c in rows[1] if c):
            blocks.append(("table", [rows[0]] + rows[2:]))
        else:
            blocks.append(("p", " ".join(table)))
        table.clear()

    for raw in content.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if code is not None:
            if line.strip().startswith("```"):
                blocks.append(("code", "\n".join(code)))
                code = None
            else:
                code.append(raw)
            continue
        s = line.strip()
        if s.startswith("|"):
            flush()
            table.append(s)
            continue
        flush_table()
        if s.startswith("```"):
            flush()
            code = []
        elif not s:
            flush()
        elif m := re.match(r"^(#{1,4})\s+(.*)", s):
            flush()
            blocks.append(("h", (len(m.group(1)), m.group(2))))
        elif re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            flush()
            blocks.append(("hr", None))
        elif m := re.match(r"^[-*•]\s+(.*)", s):
            flush()
            blocks.append(("ul", m.group(1)))
        elif m := re.match(r"^(\d+)[.)]\s+(.*)", s):
            flush()
            blocks.append(("ol", (m.group(1), m.group(2))))
        else:
            para.append(s)
    flush_table()
    if code is not None:
        blocks.append(("code", "\n".join(code)))
    flush()
    return blocks


def inline_runs(text):
    """Yield (text, bold, italic, mono) segments."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            yield part[2:-2], True, False, False
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            yield part[1:-1], False, False, True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            yield part[1:-1], False, True, False
        else:
            yield part, False, False, False


# ---------- DOCX ----------

def _docx_inline(paragraph, text):
    for chunk, bold, italic, mono in inline_runs(text):
        run = paragraph.add_run(chunk)
        run.bold, run.italic = bold, italic
        if mono:
            run.font.name = "Consolas"


def to_docx(title, content):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, 0)
    for kind, data in parse_blocks(content):
        if kind == "h":
            level, text = data
            doc.add_heading(text.strip("*"), min(level, 4))
        elif kind == "p":
            p = doc.add_paragraph()
            for n, line in enumerate(data.split("\n")):
                if n:
                    p.add_run().add_break()
                _docx_inline(p, line)
        elif kind == "ul":
            _docx_inline(doc.add_paragraph(style="List Bullet"), data)
        elif kind == "ol":
            _docx_inline(doc.add_paragraph(style="List Number"), data[1])
        elif kind == "code":
            run = doc.add_paragraph().add_run(data)
            run.font.name, run.font.size = "Consolas", Pt(9.5)
        elif kind == "table":
            cols = max(len(r) for r in data)
            table = doc.add_table(rows=0, cols=cols)
            table.style = "Table Grid"
            for n, row in enumerate(data):
                cells = table.add_row().cells
                for i in range(cols):
                    p = cells[i].paragraphs[0]
                    _docx_inline(p, row[i] if i < len(row) else "")
                    if n == 0:
                        for run in p.runs:
                            run.bold = True
            doc.add_paragraph()
        elif kind == "hr":
            doc.add_paragraph("―" * 30)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------- PDF ----------

_FONTS = None


def _pdf_fonts():
    """Prefer a system TTF (wider glyph coverage) over the built-in Helvetica."""
    global _FONTS
    if _FONTS:
        return _FONTS
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.fonts import addMapping

    fonts_dir = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    candidates = [
        ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"),
        ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf"),
    ]
    search = [fonts_dir, "/usr/share/fonts/truetype/dejavu", "/Library/Fonts"]
    for files in candidates:
        for folder in search:
            paths = [os.path.join(folder, f) for f in files]
            if all(os.path.exists(p) for p in paths[:2]):
                names = ["DocSans", "DocSans-Bold", "DocSans-Italic", "DocSans-BoldItalic"]
                for i, (name, path) in enumerate(zip(names, paths)):
                    pdfmetrics.registerFont(TTFont(name, path if os.path.exists(path) else paths[i % 2]))
                addMapping("DocSans", 0, 0, "DocSans")
                addMapping("DocSans", 1, 0, "DocSans-Bold")
                addMapping("DocSans", 0, 1, "DocSans-Italic")
                addMapping("DocSans", 1, 1, "DocSans-BoldItalic")
                _FONTS = ("DocSans", "DocSans-Bold")
                return _FONTS
    _FONTS = ("Helvetica", "Helvetica-Bold")
    return _FONTS


def _pdf_markup(text):
    out = []
    for chunk, bold, italic, mono in inline_runs(text):
        chunk = chunk.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if mono:
            chunk = f'<font face="Courier">{chunk}</font>'
        if bold:
            chunk = f"<b>{chunk}</b>"
        if italic:
            chunk = f"<i>{chunk}</i>"
        out.append(chunk)
    return "".join(out)


def to_pdf(title, content):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (HRFlowable, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    regular, bold = _pdf_fonts()
    body = ParagraphStyle("body", fontName=regular, fontSize=11, leading=15.5, spaceAfter=7)
    heads = {
        0: ParagraphStyle("title", fontName=bold, fontSize=20, leading=25, spaceAfter=14),
        1: ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, spaceBefore=10, spaceAfter=6),
        2: ParagraphStyle("h2", fontName=bold, fontSize=13.5, leading=17, spaceBefore=8, spaceAfter=5),
        3: ParagraphStyle("h3", fontName=bold, fontSize=12, leading=15, spaceBefore=6, spaceAfter=4),
    }
    item = ParagraphStyle("item", parent=body, leftIndent=16, bulletIndent=4, spaceAfter=3)
    code = ParagraphStyle("code", fontName="Courier", fontSize=9, leading=12, spaceAfter=8)

    story = [Paragraph(_pdf_markup(title), heads[0])]
    for kind, data in parse_blocks(content):
        if kind == "h":
            story.append(Paragraph(_pdf_markup(data[1]), heads[min(data[0], 3)]))
        elif kind == "p":
            story.append(Paragraph("<br/>".join(_pdf_markup(line) for line in data.split("\n")), body))
        elif kind == "ul":
            story.append(Paragraph(_pdf_markup(data), item, bulletText="•"))
        elif kind == "ol":
            story.append(Paragraph(_pdf_markup(data[1]), item, bulletText=f"{data[0]}."))
        elif kind == "code":
            story.append(Preformatted(data, code))
        elif kind == "table":
            cols = max(len(r) for r in data)
            cell = ParagraphStyle("cell", parent=body, fontSize=9.5, leading=12.5, spaceAfter=0)
            head = ParagraphStyle("head", parent=cell, fontName=bold)
            rows = [[Paragraph(_pdf_markup(r[i] if i < len(r) else ""), head if n == 0 else cell) for i in range(cols)]
                    for n, r in enumerate(data)]
            grid = Table(rows, colWidths=[(A4[0] - 4.4 * cm) / cols] * cols, repeatRows=1)
            grid.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9aa9b3")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6edf1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [grid, Spacer(1, 8)]
        elif kind == "hr":
            story.append(HRFlowable(width="100%", thickness=0.6, spaceBefore=4, spaceAfter=8))
    story.append(Spacer(1, 1))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=title, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm).build(story)
    buf.seek(0)
    return buf


def to_markdown(title, content):
    return io.BytesIO(f"# {title}\n\n{content.strip()}\n".encode("utf-8"))


EXPORTERS = {
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": (to_pdf, "application/pdf"),
    "md": (to_markdown, "text/markdown"),
}


def export_document(title, content, fmt):
    if fmt not in EXPORTERS:
        raise ValueError(f"Unsupported format: {fmt}")
    if not content.strip():
        raise ValueError("Document is empty.")
    render, mime = EXPORTERS[fmt]
    safe = re.sub(r"[\s_]+", "_", re.sub(r"[^\w\- ]+", "", title).strip())[:60] or "document"
    return render(title, content), mime, f"{safe}.{fmt}"
